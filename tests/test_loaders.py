from pathlib import Path

import pytest
from docx import Document

from src.loaders import load_file


def test_txt_is_one_page(tmp_path: Path):
    file = tmp_path / "policy.txt"
    file.write_text("Returns within 30 days.", encoding="utf-8")
    assert load_file(file) == [(1, "Returns within 30 days.")]


def test_unsupported_type_is_rejected(tmp_path: Path):
    file = tmp_path / "notes.xyz"
    file.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        load_file(file)


def test_missing_file_raises(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        load_file(tmp_path / "nope.txt")


def test_docx_is_one_page(tmp_path: Path):
    file = tmp_path / "handbook.docx"
    doc = Document()
    doc.add_paragraph("Leave policy")
    doc.add_paragraph("Employees get 20 days.")
    doc.save(str(file))
    assert load_file(file) == [(1, "Leave policy\nEmployees get 20 days.")]
